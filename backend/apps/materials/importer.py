"""Read .xlsx / .csv uploads into clean rows. No Django imports — unit-tested on its own.

Finds the header row automatically (the High-Rise register has two title rows above it),
matches column names loosely ("Material / equipment item", "Item", "Material" …) and converts
numbers and dates. Problems are reported per row instead of failing the whole file.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

MAX_ROWS = 20000
MAX_BYTES = 10 * 1024 * 1024


class ImportError_(Exception):
    """Whole-file problem (wrong type, no header row…). Message is user-facing."""


def norm(s) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()


# field -> (type, aliases). First alias set wins when two columns could match.
CATALOGUE = {
    "code": ("code", ["code", "item code", "material code", "sku", "item id", "id"]),
    "name": ("text", ["name", "material name", "material", "item name", "item", "description", "material description"]),
    "category": ("text", ["category", "group", "trade", "material group"]),
    "unit": ("text", ["unit", "uom", "units"]),
    "spec": ("text", ["spec", "specification", "specifications", "grade"]),
    "hsn": ("text", ["hsn", "hsn code", "hsn sac"]),
    "gst_rate": ("number", ["gst", "gst rate", "gst %", "gst percent", "tax rate"]),
    "std_rate": ("number", ["rate", "std rate", "standard rate", "unit rate", "price", "basic rate"]),
    "reorder_level": ("number", ["reorder level", "reorder", "min stock", "minimum stock", "reorder qty"]),
}
CATALOGUE_KEY, CATALOGUE_REQUIRED = "code", ["code", "name"]

REGISTER = {
    "item_id": ("code", ["item id", "id", "item code", "code"]),
    "trade": ("text", ["trade"]),
    "system": ("text", ["system section", "system", "section"]),
    "item": ("text", ["material equipment item", "material item", "item", "material", "description"]),
    "unit": ("text", ["unit", "uom"]),
    "spec": ("text", ["specification to confirm", "specification", "spec"]),
    "applicability": ("choice", ["applicability", "applicable"]),
    "zone": ("text", ["block floor zone", "zone", "location", "block"]),
    "req_qty": ("number", ["required quantity", "quantity", "qty", "required qty"]),
    "grade": ("text", ["approved size grade rating", "grade", "size", "rating"]),
    "make": ("text", ["approved make", "make", "brand"]),
    "vendor": ("text", ["vendor", "supplier"]),
    "boq_ref": ("text", ["drawing boq reference", "boq reference", "boq ref", "drawing reference"]),
    "status": ("choice", ["procurement status", "status"]),
    "required_by": ("date", ["required by date", "required by", "need by"]),
    "remarks": ("text", ["remarks package inclusion", "remarks", "notes"]),
}
REGISTER_KEY, REGISTER_REQUIRED = "item_id", ["item_id", "item"]

APPLICABILITY = ["To review", "Required", "Included in package", "Not applicable"]
REG_STATUS = ["Not started", "Specs in review", "Specs approved", "Enquiry / quotes", "Ordered", "Partly delivered",
              "Delivered", "On hold"]
CHOICES = {"applicability": APPLICABILITY, "status": REG_STATUS}


# ---------------------------------------------------------------- reading
def read_table(filename: str, content: bytes) -> list[list]:
    if len(content) > MAX_BYTES:
        raise ImportError_("File is larger than 10 MB — split it and upload again.")
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:  # pragma: no cover
            raise ImportError_("Excel support isn’t installed on the server.")
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise ImportError_("Couldn’t open this Excel file. Save it again as .xlsx and retry.")
        ws = _pick_sheet(wb)
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append(list(r))
            if len(rows) > MAX_ROWS + 20:
                raise ImportError_(f"More than {MAX_ROWS:,} rows — split the file.")
        return rows
    if name.endswith((".csv", ".txt")):
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        rows = list(csv.reader(io.StringIO(text or "")))
        if len(rows) > MAX_ROWS + 20:
            raise ImportError_(f"More than {MAX_ROWS:,} rows — split the file.")
        return rows
    if name.endswith(".xls"):
        raise ImportError_("Old .xls files aren’t supported — open it in Excel and save as .xlsx.")
    raise ImportError_("Upload an .xlsx or .csv file.")


def _pick_sheet(wb):
    """Prefer a sheet called like the register / materials; otherwise the first one."""
    for ws in wb.worksheets:
        if any(k in norm(ws.title) for k in ("register", "material", "catalogue", "items")):
            return ws
    return wb.worksheets[0]


# ---------------------------------------------------------------- header matching
def map_header(row: list, spec: dict) -> dict[str, int]:
    """field -> column index. Exact alias match first, then 'header starts with alias'."""
    cells = [norm(c) for c in row]
    out: dict[str, int] = {}
    used: set[int] = set()
    for exact in (True, False):
        for field, (_, aliases) in spec.items():
            if field in out:
                continue
            for alias in aliases:
                hit = next((ix for ix, c in enumerate(cells) if ix not in used and c and
                            (c == alias if exact else (c.startswith(alias + " ") or alias.startswith(c + " ")))), None)
                if hit is not None:
                    out[field] = hit
                    used.add(hit)
                    break
    return out


def find_header(rows: list[list], spec: dict, required: list[str]) -> tuple[int, dict]:
    for ix, row in enumerate(rows[:25]):
        m = map_header(row, spec)
        if all(f in m for f in required):
            return ix, m
    need = ", ".join(spec[f][1][0].title() for f in required)
    raise ImportError_(f"Couldn’t find the header row. The file needs columns for: {need}.")


# ---------------------------------------------------------------- values
def to_number(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool):
        return Decimal(str(v))
    s = re.sub(r"[₹,\s%]", "", str(v)).replace("Rs.", "").replace("Rs", "")
    if s in ("", "-", "—"):
        return None
    try:
        return Decimal(s)
    except InvalidOperation:
        raise ValueError(f"“{v}” isn’t a number")


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y", "%b %Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"“{v}” isn’t a date (use DD-MM-YYYY)")


def to_text(v, limit=255):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()[:limit]


def to_code(v):
    s = to_text(v, 80).upper()
    return re.sub(r"\s+", "-", s)


def to_choice(v, field):
    s = to_text(v)
    if not s:
        return None
    for c in CHOICES[field]:
        if norm(c) == norm(s):
            return c
    raise ValueError(f"“{s}” isn’t a valid {field.replace('_', ' ')} ({', '.join(CHOICES[field])})")


def parse(rows: list[list], spec: dict, key: str, required: list[str]) -> dict:
    """-> {"columns": {field: header text}, "rows": [{"line", "data"}], "errors": [{"line", "message"}]}"""
    hix, cols = find_header(rows, spec, required)
    header = rows[hix]
    good, errors, seen = [], [], {}
    for ix in range(hix + 1, len(rows)):
        raw = rows[ix]
        line = ix + 1
        if not raw or all(c in (None, "") or str(c).strip() == "" for c in raw):
            continue
        data, problems = {}, []
        for field, col in cols.items():
            v = raw[col] if col < len(raw) else None
            kind = spec[field][0]
            try:
                if kind == "number":
                    data[field] = to_number(v)
                elif kind == "date":
                    data[field] = to_date(v)
                elif kind == "code":
                    data[field] = to_code(v)
                elif kind == "choice":
                    data[field] = to_choice(v, field)
                else:
                    data[field] = to_text(v, 2000 if field in ("spec", "remarks") else 255)
            except ValueError as e:
                problems.append(f"{header[col]}: {e}")
        for f in required:
            if not data.get(f):
                problems.append(f"{spec[f][1][0].title()} is empty")
        k = data.get(key)
        if k and k in seen:
            problems.append(f"Duplicate {key.replace('_', ' ')} {k} (also on line {seen[k]})")
        if problems:
            errors.append({"line": line, "message": "; ".join(problems)})
            continue
        seen[k] = line
        good.append({"line": line, "data": {f: v for f, v in data.items() if v is not None}})
    return {"columns": {f: to_text(header[c]) for f, c in cols.items()}, "rows": good, "errors": errors}


def parse_catalogue(filename, content):
    return parse(read_table(filename, content), CATALOGUE, CATALOGUE_KEY, CATALOGUE_REQUIRED)


def parse_register(filename, content):
    return parse(read_table(filename, content), REGISTER, REGISTER_KEY, REGISTER_REQUIRED)
