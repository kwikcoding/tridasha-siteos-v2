"""Cross-module report for a period, filtered by site and project. Each section appears only if the person's role
can see that module — financial sections (purchase money, wages, expenses) follow the same rule."""
from datetime import timedelta
from decimal import Decimal

from django.db.models import Count, DecimalField, ExpressionWrapper, F, Q, Sum
from django.db.models.functions import Coalesce
from django.utils import timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Font
from rest_framework import permissions
from rest_framework.decorators import api_view, permission_classes
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response

from apps.accounts.permissions import module_permission
from apps.core.scoped import user_site_ids

ZERO = Decimal("0")


def _period(p):
    today = timezone.localdate()
    start = parse_date(p.get("from") or "") or today.replace(day=1)
    end = parse_date(p.get("to") or "") or today
    if end < start:
        raise ValidationError({"to": "End date is before the start date."})
    if (end - start).days > 1100:
        raise ValidationError({"from": "Choose a period of up to three years."})
    return start, end


def _scope(qs, user, p, site_path="site", project=True):
    ids = user_site_ids(user)
    if ids is not None:
        qs = qs.filter(**{f"{site_path}__in": ids})
    if p.get("site"):
        qs = qs.filter(**{f"{site_path}_id": p["site"]})
    if project and p.get("project"):
        qs = qs.filter(project_id=p["project"])
    return qs


def _f(v):
    return float(v or 0)


def build(user, p):
    start, end = _period(p)
    out = {"from": start, "to": end, "sections": []}
    today = timezone.localdate()

    if user.can_view("indents"):
        from apps.indents import workflow as wf
        from apps.indents.models import Indent
        qs = _scope(Indent.objects.all(), user, p)
        raised = qs.filter(created_at__date__range=(start, end))
        money = qs.filter(pclosed_at__date__range=(start, end)) | qs.filter(pstage__in=["invoice", "payauth", "payment"])
        m = money.aggregate(inv=Sum("inv_amount"), gst=Sum("inv_gst"), auth=Sum("pay_auth_amount"), paid=Sum("paid_total"))
        out["indents"] = {
            "raised": raised.count(),
            "closed": qs.filter(closed_at__date__range=(start, end)).count(),
            "open": qs.exclude(status__in=list(wf.FINAL)).count(),
            "rejected_or_cancelled": raised.filter(status__in=["Rejected", "Cancelled"]).count(),
            "by_status": list(qs.exclude(status__in=list(wf.FINAL)).values("status").annotate(n=Count("id")).order_by("-n")),
            "invoiced": _f(m["inv"]) + _f(m["gst"]), "authorised": _f(m["auth"]), "paid": _f(m["paid"]),
            "to_pay": max(_f(m["auth"]) - _f(m["paid"]), 0),
        }
        out["sections"].append("indents")

    if user.can_view("materials"):
        from apps.materials.services import low_stock_count
        out["materials"] = {"low_stock": low_stock_count(user)}
        out["sections"].append("materials")

    if user.can_view("dpr"):
        from apps.sitework.models import DailyReport
        qs = _scope(DailyReport.objects.all(), user, p).filter(date__range=(start, end))
        days = (end - start).days + 1
        out["dpr"] = {"reports": qs.count(), "approved": qs.filter(status="approved").count(),
                      "waiting_approval": qs.filter(status="submitted").count(),
                      "manpower": sum(r.manpower_total for r in qs.only("manpower")),
                      "days_in_period": days}
        out["sections"].append("dpr")

    if user.can_view("labour"):
        from apps.sitework.models import Attendance
        qs = _scope(Attendance.objects.all(), user, p).filter(date__range=(start, end))
        a = qs.aggregate(present=Count("id", filter=Q(status="present")), half=Count("id", filter=Q(status="half")),
                         ot=Sum("ot_hours"), wages=Sum("wage"))
        out["labour"] = {"man_days": a["present"] + 0.5 * a["half"], "ot_hours": _f(a["ot"]), "wages": _f(a["wages"]),
                         "unapproved_days": qs.filter(approved=False).values("site", "date").distinct().count(),
                         "by_contractor": [{"contractor": r["worker__contractor"] or "Departmental",
                                            "man_days": r["p"] + 0.5 * r["h"], "wages": _f(r["w"])}
                                           for r in qs.values("worker__contractor").annotate(
                                               p=Count("id", filter=Q(status="present")),
                                               h=Count("id", filter=Q(status="half")), w=Sum("wage"))
                                           .order_by("worker__contractor")]}
        out["sections"].append("labour")

    if user.can_view("expenses"):
        from apps.finance.models import Expense
        qs = _scope(Expense.objects.all(), user, p).filter(date__range=(start, end)).exclude(status="rejected")
        t = qs.aggregate(amount=Sum("amount"), gst=Sum("gst"), approved=Sum("approved_amount"), paid=Sum("paid_total"))
        out["expenses"] = {
            "bills": qs.count(), "billed": _f(t["amount"]) + _f(t["gst"]), "approved": _f(t["approved"]),
            "paid": _f(t["paid"]), "outstanding": max(_f(t["approved"]) - _f(t["paid"]), 0),
            "waiting_approval": qs.filter(status="submitted").count(),
            "by_category": [{"category": r["category"], "bills": r["n"], "total": _f(r["a"]) + _f(r["g"]),
                             "paid": _f(r["p"])}
                            for r in qs.values("category").annotate(n=Count("id"), a=Sum("amount"), g=Sum("gst"),
                                                                    p=Sum("paid_total")).order_by("category")],
            "by_department": [{"department": r["created_by__department__name"] or "—", "bills": r["n"],
                               "total": _f(r["a"]) + _f(r["g"])}
                              for r in qs.values("created_by__department__name").annotate(
                                  n=Count("id"), a=Sum("amount"), g=Sum("gst")).order_by("created_by__department__name")],
        }
        out["sections"].append("expenses")

    if user.can_view("tasks"):
        from apps.sitework.models import Task
        qs = _scope(Task.objects.all(), user, p)
        open_ = qs.exclude(status__in=Task.CLOSED)
        out["tasks"] = {"open": open_.count(), "overdue": open_.filter(due_date__lt=today).count(),
                        "done_in_period": qs.filter(completed_at__date__range=(start, end)).count(),
                        "due_in_period": qs.filter(due_date__range=(start, end)).count(),
                        "by_department": [{"department": r["assignee__department__name"] or "Unassigned", "open": r["n"],
                                           "overdue": r["o"]}
                                          for r in open_.values("assignee__department__name").annotate(
                                              n=Count("id"), o=Count("id", filter=Q(due_date__lt=today)))
                                          .order_by("assignee__department__name")]}
        out["sections"].append("tasks")

    if user.can_view("quality"):
        from apps.compliance.models import Observation
        qs = _scope(Observation.objects.all(), user, p)
        period = qs.filter(date__range=(start, end))
        out["quality"] = {"raised": period.count(), "closed": qs.filter(closed_at__date__range=(start, end)).count(),
                          "open": qs.exclude(status="closed").count(),
                          "overdue": qs.exclude(status="closed").filter(due_date__lt=today).count(),
                          "incidents": period.filter(kind="incident").count(),
                          "by_kind": list(period.values("kind").annotate(n=Count("id")).order_by("kind"))}
        out["sections"].append("quality")

    if user.can_view("equipment"):
        from apps.plant.models import Equipment, MaintenanceLog
        qs = Equipment.objects.exclude(status="disposed")
        ids = user_site_ids(user)
        if ids is not None:
            qs = qs.filter(Q(site__in=ids) | Q(site__isnull=True))
        if p.get("site"):
            qs = qs.filter(site_id=p["site"])
        logs = MaintenanceLog.objects.filter(equipment__in=qs, date__range=(start, end)).aggregate(
            c=Sum("cost"), d=Sum("downtime_hours"))
        out["equipment"] = {"total": qs.count(), "in_use": qs.filter(status="in_use").count(),
                            "out_of_service": qs.filter(status__in=["maintenance", "breakdown"]).count(),
                            "service_due": qs.filter(next_service_due__lte=today).count(),
                            "maintenance_cost": _f(logs["c"]), "downtime_hours": _f(logs["d"])}
        out["sections"].append("equipment")

    out["projects"] = project_rows(user, p, start, end)
    return out


def project_rows(user, p, start, end):
    from apps.sites.models import Project
    qs = _scope(Project.objects.select_related("site"), user, p, project=False)
    if p.get("project"):
        qs = qs.filter(pk=p["project"])
    today = timezone.localdate()
    rows = []
    for pr in qs:
        r = {"id": pr.id, "site": pr.site.name, "code": pr.code, "name": pr.name, "status": pr.get_status_display(),
             "budget": _f(pr.budget)}
        if user.can_view("indents"):
            from apps.indents import workflow as wf
            r["open_indents"] = pr.indents.exclude(status__in=list(wf.FINAL)).count()
        if user.can_view("expenses"):
            t = pr.expenses.exclude(status="rejected").aggregate(a=Sum("approved_amount"), p=Sum("paid_total"))
            r["approved_spend"], r["paid"] = _f(t["a"]), _f(t["p"])
            r["budget_used_pct"] = round(100 * _f(t["a"]) / _f(pr.budget), 1) if pr.budget else None
        if user.can_view("dpr"):
            r["reports"] = pr.daily_reports.filter(date__range=(start, end)).count()
        if user.can_view("tasks"):
            r["overdue_tasks"] = pr.tasks.exclude(status__in=["done", "cancelled"]).filter(due_date__lt=today).count()
        if user.can_view("quality"):
            r["open_observations"] = pr.observations.exclude(status="closed").count()
        rows.append(r)
    return rows


@api_view(["GET"])
@permission_classes([permissions.IsAuthenticated, module_permission("reports")])
def summary(request):
    return Response(build(request.user, request.query_params))


@api_view(["GET"])
@permission_classes([permissions.IsAuthenticated, module_permission("reports")])
def export(request):
    from django.http import HttpResponse
    if not request.user.has_code("reports.export"):
        from rest_framework.exceptions import PermissionDenied
        raise PermissionDenied("Your role doesn’t allow this (reports.export).")
    data = build(request.user, request.query_params)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    from apps.core.models import CompanyProfile
    co = CompanyProfile.load()
    ws.append([f"{co.legal_name or co.name} — SiteOS report {data['from']:%d %b %Y} – {data['to']:%d %b %Y}"])
    if co.gstin:
        ws.append([f"GSTIN {co.gstin}"])
    ws["A1"].font = Font(bold=True, size=13)
    for sec in data["sections"]:
        ws.append([])
        ws.append([sec.title()])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for k, v in data[sec].items():
            if not isinstance(v, list):
                ws.append([k.replace("_", " ").capitalize(), v])
        for k, v in data[sec].items():
            if isinstance(v, list) and v:
                sheet = wb.create_sheet(f"{sec} {k}"[:31])
                sheet.append(list(v[0].keys()))
                for c in sheet[1]:
                    c.font = Font(bold=True)
                for row in v:
                    sheet.append([x if not isinstance(x, Decimal) else float(x) for x in row.values()])
    if data["projects"]:
        sheet = wb.create_sheet("Projects")
        keys = list(data["projects"][0].keys())
        sheet.append(keys)
        for c in sheet[1]:
            c.font = Font(bold=True)
        for r in data["projects"]:
            sheet.append([r.get(k) for k in keys])
    ws.column_dimensions["A"].width = 32
    resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="siteos-report-{data["from"]}-{data["to"]}.xlsx"'
    wb.save(resp)
    from apps.audit.log import record
    record("reports.export", actor=request.user, summary=f"Downloaded report {data['from']} – {data['to']}",
           request=request)
    return resp
