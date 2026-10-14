"""Security and data-isolation tests: CSRF, headers, throttling, mass assignment, files, and site isolation for
every module (someone at another site gets 404 on every record, action, file and export row)."""
import shutil
import tempfile
from datetime import timedelta
from decimal import Decimal

from django.contrib.contenttypes.models import ContentType
from django.core.cache import cache
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.utils import timezone
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.compliance.models import Document, DocumentRevision, Observation
from apps.core.models import Attachment
from apps.finance.models import Expense
from apps.indents.models import Indent
from apps.materials.models import Material, StockEntry
from apps.plant.models import Equipment
from apps.sites.models import Project, Site
from apps.sitework.models import Attendance, DailyReport, Task, Worker

PW = "Str0ng-pass-123"
MEDIA = tempfile.mkdtemp(prefix="siteos-sec-media-")
TODAY = timezone.localdate()


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, MEDIA_ROOT=MEDIA)
class SecurityBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed", verbosity=0)
        cls.a = Site.objects.create(code="TIM", name="Tridasha Imperial")
        cls.b = Site.objects.create(code="VAR", name="Varuna")
        cls.admin = User.objects.create_superuser("9876543210", PW, name="Admin")
        # a person with every module at site B only (management role, but limited to site B)
        role = Role.objects.get(key="pm")
        role.all_sites = False
        role.save()
        cls.outsider = User.objects.create_user("9700000001", PW, name="Outsider", role=role)
        cls.outsider.sites.add(cls.b)
        cls.insider = User.objects.create_user("9700000002", PW, name="Insider", role=role)
        cls.insider.sites.add(cls.a)
        u = cls.insider
        p = Project.objects.create(site=cls.a, code="TA", name="Tower A")
        m = Material.objects.create(code="CEM", name="Cement", unit="bag")
        StockEntry.objects.create(site=cls.a, material=m, kind="IN", qty=10, date=TODAY, by=u)
        w = Worker.objects.create(site=cls.a, name="W", daily_wage=500, created_by=u)
        cls.objs = {
            "projects": p,
            "tasks": Task.objects.create(site=cls.a, project=p, title="T", created_by=u),
            "dpr": DailyReport.objects.create(site=cls.a, date=TODAY, summary="S", created_by=u, status="submitted"),
            "workers": w,
            "attendance": Attendance.objects.create(worker=w, site=cls.a, date=TODAY, status="present", created_by=u),
            "expenses": Expense.objects.create(site=cls.a, date=TODAY, category="Fuel", description="D", payee="P",
                                               amount=Decimal("100"), status="submitted", created_by=u),
            "observations": Observation.objects.create(site=cls.a, kind="safety", ref_no="SO-1", title="O", created_by=u),
            "documents": Document.objects.create(site=cls.a, category="drawing", number="D-1", title="D", created_by=u),
            "equipment": Equipment.objects.create(code="EQ-1", name="E", site=cls.a, status="in_use", created_by=u),
            "indents": Indent.objects.create(number="IND-X-001", site=cls.a, item="Sand", qty=1, raised_by=u,
                                             stage="stock", status="Requested"),
        }
        r = DocumentRevision(document=cls.objs["documents"], rev="R0", file_name="d.pdf", uploaded_by=u)
        r.file.save("d.pdf", ContentFile(b"%PDF secret"), save=False)
        r.save()
        cls.rev = r
        att = Attachment(content_type=ContentType.objects.get_for_model(Observation), object_id=cls.objs["observations"].pk,
                         name="secret.jpg", uploaded_by=u)
        att.file.save("secret.jpg", ContentFile(b"secret"), save=False)
        att.save()
        cls.att = att

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    def setUp(self):
        cache.clear()

    def c(self, user):
        c = APIClient()
        c.force_authenticate(user)
        return c


class SiteIsolation(SecurityBase):
    ACTIONS = {"tasks": ["progress"], "dpr": ["approve", "send-back", "submit"], "expenses": ["approve", "reject", "pay", "submit"],
               "observations": ["close", "start", "reopen"], "documents": ["archive"],
               "indents": ["act"]}

    def test_every_module_record_is_invisible_from_another_site(self):
        out, ins = self.c(self.outsider), self.c(self.insider)
        leaks = []
        for ep, obj in self.objs.items():
            url = f"/api/{ep}/{obj.pk}/"
            if ins.get(url).status_code != 200:
                leaks.append(f"insider can't read {url}")
            for method in ("get", "patch", "delete"):
                r = getattr(out, method)(url, {}, format="json")
                # 403 on DELETE comes from the permission layer before any lookup (no delete right) — reveals nothing
                if r.status_code not in ((404, 405, 403) if method == "delete" else (404, 405)):
                    leaks.append(f"{method.upper()} {url} → {r.status_code}")
            for a in self.ACTIONS.get(ep, []):
                r = out.post(f"{url}{a}/", {"note": "x", "action": "approve", "date": TODAY, "amount": "1"}, format="json")
                if r.status_code != 404:
                    leaks.append(f"POST {url}{a}/ → {r.status_code}")
            lst = out.get(f"/api/{ep}/").json()
            rows = lst["results"] if isinstance(lst, dict) else lst
            if any(x.get("id") == obj.pk for x in rows if isinstance(x, dict)) and ep != "equipment":
                leaks.append(f"{ep} list shows site A record")
            r = out.get(f"/api/{ep}/", {"site": self.a.pk})
            body = r.json()
            rows = body["results"] if isinstance(body, dict) and "results" in body else body
            if r.status_code == 200 and isinstance(rows, list) and len(rows) and ep not in ("equipment",):
                leaks.append(f"{ep}?site=A returned {len(rows)} rows")
        # files
        for url in [f"/api/observations/{self.objs['observations'].pk}/attachments/{self.att.pk}/",
                    f"/api/documents/{self.objs['documents'].pk}/revisions/{self.rev.pk}/download/"]:
            if out.get(url).status_code != 404:
                leaks.append(f"file {url} reachable")
            if ins.get(url).status_code != 200:
                leaks.append(f"insider can't fetch {url}")
        # exports only contain own-site rows
        from openpyxl import load_workbook
        import io
        for ep in ("tasks", "dpr", "expenses", "observations", "documents", "attendance", "workers"):
            r = out.get(f"/api/{ep}/export/")
            ws = load_workbook(io.BytesIO(r.content)).active
            if ws.max_row > 1:
                leaks.append(f"{ep} export has {ws.max_row - 1} site A rows")
        # reports and dashboard counts don't include site A
        rep = out.get("/api/reports/summary/").json()
        if rep.get("expenses", {}).get("bills"):
            leaks.append("report counts site A bills")
        self.assertEqual(leaks, [], "\n".join(leaks))

    def test_equipment_at_another_site_is_hidden_but_free_equipment_is_shared(self):
        free = Equipment.objects.create(code="EQ-FREE", name="Free mixer", status="available", created_by=self.insider)
        codes = [e["code"] for e in self.c(self.outsider).get("/api/equipment/").json()["results"]]
        self.assertIn(free.code, codes)
        self.assertNotIn("EQ-1", codes)

    def test_cannot_create_records_for_a_site_you_are_not_on(self):
        out = self.c(self.outsider)
        p = self.objs["projects"]
        bodies = {"tasks": {"site": self.a.pk, "title": "x"}, "dpr": {"site": self.a.pk, "date": TODAY, "summary": "x"},
                  "workers": {"site": self.a.pk, "name": "x"}, "expenses": {"site": self.a.pk, "date": TODAY, "category": "Fuel",
                                                                           "description": "x", "payee": "x", "amount": "1"},
                  "observations": {"site": self.a.pk, "kind": "safety", "title": "x"},
                  "documents": {"site": self.a.pk, "category": "drawing", "number": "Z", "title": "x"},
                  "projects": {"site": self.a.pk, "code": "ZZ", "name": "x"}, "indents": {"site": self.a.pk, "item": "x", "qty": 1}}
        for ep, body in bodies.items():
            r = out.post(f"/api/{ep}/", body, format="json")
            self.assertEqual(r.status_code, 400, (ep, r.content[:200]))
        # nor attach a record at site B to a project at site A
        r = out.post("/api/tasks/", {"site": self.b.pk, "project": p.pk, "title": "x"}, format="json")
        self.assertEqual(r.status_code, 400)
        r = out.post("/api/attendance/mark/", {"site": self.a.pk, "date": TODAY, "entries": []}, format="json")
        self.assertEqual(r.status_code, 400)


class WebSecurity(SecurityBase):
    def test_csrf_is_enforced_on_session_requests(self):
        c = APIClient(enforce_csrf_checks=True)
        self.assertEqual(c.post("/api/auth/login/", {"mobile": self.insider.mobile, "password": PW}, format="json").status_code, 403)
        c.get("/api/auth/csrf/")
        token = c.cookies["csrftoken"].value
        r = c.post("/api/auth/login/", {"mobile": self.insider.mobile, "password": PW}, format="json", HTTP_X_CSRFTOKEN=token)
        self.assertEqual(r.status_code, 200)
        self.assertTrue(c.cookies["sessionid"]["httponly"])
        token = c.cookies["csrftoken"].value                     # rotated at login
        self.assertEqual(c.post("/api/tasks/", {"site": self.a.pk, "title": "x"}, format="json").status_code, 403)
        r = c.post("/api/tasks/", {"site": self.a.pk, "title": "x"}, format="json", HTTP_X_CSRFTOKEN=token)
        self.assertEqual(r.status_code, 201)

    def test_security_headers(self):
        r = self.c(self.insider).get("/api/auth/me/")
        self.assertEqual(r["X-Frame-Options"], "DENY")
        self.assertEqual(r["X-Content-Type-Options"], "nosniff")
        self.assertIn("Referrer-Policy", r)

    def test_login_is_rate_limited(self):
        c = APIClient()
        codes = [c.post("/api/auth/login/", {"mobile": "9000000000", "password": "wrong"}, format="json").status_code
                 for _ in range(12)]
        self.assertIn(429, codes)
        self.assertEqual(codes[0], 400)

    def test_signed_out_gets_nothing_from_any_module(self):
        c = APIClient()
        for ep in ["projects", "tasks", "dpr", "workers", "attendance", "expenses", "observations", "documents",
                   "equipment", "indents", "materials", "stock", "users", "audit", "people", "company", "reports/summary"]:
            self.assertEqual(c.get(f"/api/{ep}/").status_code, 403, ep)

    def test_status_and_money_fields_cannot_be_set_directly(self):
        ins = self.c(self.insider)
        e = self.objs["expenses"]
        ins.patch(f"/api/expenses/{e.pk}/", {"status": "paid", "approved_amount": "999", "paid_total": "999"}, format="json")
        e.refresh_from_db()
        self.assertEqual((e.status, e.approved_amount, e.paid_total), ("submitted", None, Decimal("0")))
        d = self.objs["dpr"]
        ins.patch(f"/api/dpr/{d.pk}/", {"status": "approved"}, format="json")
        d.refresh_from_db()
        self.assertEqual(d.status, "submitted")
        q = self.objs["equipment"]
        ins.patch(f"/api/equipment/{q.pk}/", {"status": "disposed", "site": self.b.pk}, format="json")
        q.refresh_from_db()
        self.assertEqual((q.status, q.site_id), ("in_use", self.a.pk))
        o = self.objs["observations"]
        ins.patch(f"/api/observations/{o.pk}/", {"status": "closed", "ref_no": "HACK"}, format="json")
        o.refresh_from_db()
        self.assertEqual((o.status, o.ref_no), ("open", "SO-1"))
        i = self.objs["indents"]
        self.c(self.admin).patch(f"/api/indents/{i.pk}/", {"status": "Closed", "stage": "closed", "paid_total": "5"}, format="json")
        i.refresh_from_db()
        self.assertEqual((i.status, i.stage, i.paid_total), ("Requested", "stock", None))

    def test_uploaded_file_names_cannot_escape_the_folder_and_media_is_not_public(self):
        o = self.objs["observations"]
        f = SimpleUploadedFile("../../../etc/passwd.jpg", b"x", content_type="image/jpeg")
        r = self.c(self.insider).post(f"/api/observations/{o.pk}/attach/", {"file": f}, format="multipart")
        self.assertEqual(r.status_code, 201, r.content)
        a = Attachment.objects.get(pk=r.json()["id"])
        self.assertTrue(a.file.name.startswith("uploads/observation/"))
        self.assertNotIn("..", a.file.name)
        self.assertEqual(APIClient().get("/media/" + a.file.name).status_code, 404)     # only through the API
        self.assertEqual(APIClient().get(f"/api/observations/{o.pk}/attachments/{a.pk}/").status_code, 403)

    def test_audit_log_is_append_only(self):
        from apps.audit.models import AuditLog
        n = AuditLog.objects.count()
        self.c(self.insider).post("/api/tasks/", {"site": self.a.pk, "title": "audited"}, format="json")
        self.assertEqual(AuditLog.objects.count(), n + 1)
        last = AuditLog.objects.latest("id")
        adm = self.c(self.admin)
        for method in ("patch", "put", "delete"):
            self.assertIn(getattr(adm, method)(f"/api/audit/{last.pk}/", {"summary": "x"}, format="json").status_code,
                          (403, 404, 405))      # there is no route that changes an entry
        self.assertEqual(AuditLog.objects.count(), n + 1)

    def test_search_input_is_treated_as_text(self):
        for ep in ["tasks", "expenses", "observations", "documents", "indents", "materials"]:
            r = self.c(self.insider).get(f"/api/{ep}/", {"search": "' OR 1=1 --"})
            self.assertEqual(r.status_code, 200, ep)
            self.assertEqual(r.json()["count"], 0, ep)
